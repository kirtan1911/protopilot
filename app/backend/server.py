from fastapi import FastAPI, APIRouter, HTTPException, Depends, UploadFile, File, Form, status
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import FileResponse, StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
import json
import uuid
import re
import bcrypt
import jwt
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Any, Dict
from datetime import datetime, timezone, timedelta
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor
from google import genai
from google.genai import types as genai_types
import whisper

# Configure logging early so it's available everywhere
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

UPLOAD_DIR = ROOT_DIR / "uploads"
EXPORT_DIR = ROOT_DIR / "exports"
UPLOAD_DIR.mkdir(exist_ok=True)
EXPORT_DIR.mkdir(exist_ok=True)

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

GEMINI_API_KEY = os.environ['GEMINI_API_KEY']
JWT_SECRET = os.environ.get('JWT_SECRET', 'change-me')
JWT_ALG = os.environ.get('JWT_ALGORITHM', 'HS256')

# Configure Gemini (new google.genai SDK)
gemini_client = genai.Client(api_key=GEMINI_API_KEY)
GEMINI_MODEL_NAME = os.environ.get('GEMINI_MODEL', 'gemini-2.5-flash')

# Load Whisper model once at startup (use "base" for speed; use "small"/"medium" for better accuracy)
WHISPER_MODEL_SIZE = os.environ.get('WHISPER_MODEL_SIZE', 'base')
logger.info(f"Loading Whisper model: {WHISPER_MODEL_SIZE} ...")
whisper_model = whisper.load_model(WHISPER_MODEL_SIZE)
logger.info("Whisper model loaded.")

app = FastAPI(title="ProtoPilot API")
api_router = APIRouter(prefix="/api")
security = HTTPBearer()

# ---------------- Models ----------------
def now_iso():
    return datetime.now(timezone.utc).isoformat()

class UserSignup(BaseModel):
    name: str
    email: EmailStr
    password: str

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class UserOut(BaseModel):
    id: str
    name: str
    email: str

class ProjectCreate(BaseModel):
    name: str
    description: Optional[str] = ""

class ProjectUpdate(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    transcript: Optional[str] = None
    requirements: Optional[Dict[str, Any]] = None
    schema_spec: Optional[Dict[str, Any]] = None
    api_spec: Optional[Dict[str, Any]] = None
    wireframe_html: Optional[str] = None

class TranscriptUpdate(BaseModel):
    transcript: str

class ExtractRequest(BaseModel):
    transcript: str

# ---------------- Auth helpers ----------------
def hash_password(pw: str) -> str:
    return bcrypt.hashpw(pw.encode(), bcrypt.gensalt()).decode()

def verify_password(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode(), hashed.encode())
    except Exception:
        return False

def create_token(user_id: str) -> str:
    payload = {"sub": user_id, "exp": datetime.now(timezone.utc) + timedelta(days=7)}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALG)

async def get_current_user(creds: HTTPAuthorizationCredentials = Depends(security)) -> dict:
    try:
        payload = jwt.decode(creds.credentials, JWT_SECRET, algorithms=[JWT_ALG])
        user_id = payload.get("sub")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")
    user = await db.users.find_one({"id": user_id}, {"_id": 0, "password": 0})
    if not user:
        raise HTTPException(status_code=401, detail="User not found")
    return user

# ---------------- LLM helpers (Gemini) ----------------
def extract_json(raw: str) -> Any:
    """Extract first JSON object/array from LLM output."""
    raw = raw.strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?", "", raw).rstrip("`").strip()
    # find first { or [
    start = None
    for i, ch in enumerate(raw):
        if ch in "{[":
            start = i
            break
    if start is None:
        raise ValueError("No JSON found in LLM output")
    depth = 0
    end = None
    open_ch = raw[start]
    close_ch = "}" if open_ch == "{" else "]"
    in_str = False
    esc = False
    for i in range(start, len(raw)):
        c = raw[i]
        if esc:
            esc = False
            continue
        if c == "\\":
            esc = True
            continue
        if c == '"':
            in_str = not in_str
            continue
        if in_str:
            continue
        if c == open_ch:
            depth += 1
        elif c == close_ch:
            depth -= 1
            if depth == 0:
                end = i + 1
                break
    if end is None:
        raise ValueError("Unbalanced JSON in LLM output")
    return json.loads(raw[start:end])

async def gemini_json(system: str, user_prompt: str) -> Any:
    """Call Gemini with a system+user prompt and parse JSON from the response."""
    response = await gemini_client.aio.models.generate_content(
        model=GEMINI_MODEL_NAME,
        contents=user_prompt,
        config=genai_types.GenerateContentConfig(
            system_instruction=system,
        ),
    )
    text = response.text
    return extract_json(text)

async def gemini_text(system: str, user_prompt: str) -> str:
    """Call Gemini with a system+user prompt and return raw text."""
    response = await gemini_client.aio.models.generate_content(
        model=GEMINI_MODEL_NAME,
        contents=user_prompt,
        config=genai_types.GenerateContentConfig(
            system_instruction=system,
        ),
    )
    return response.text

# ---------------- Auth Routes ----------------
@api_router.post("/auth/signup")
async def signup(data: UserSignup):
    existing = await db.users.find_one({"email": data.email.lower()})
    if existing:
        raise HTTPException(status_code=400, detail="Email already registered")
    user_id = str(uuid.uuid4())
    doc = {
        "id": user_id,
        "name": data.name,
        "email": data.email.lower(),
        "password": hash_password(data.password),
        "created_at": now_iso(),
    }
    await db.users.insert_one(doc)
    token = create_token(user_id)
    return {"token": token, "user": {"id": user_id, "name": data.name, "email": data.email.lower()}}

@api_router.post("/auth/login")
async def login(data: UserLogin):
    user = await db.users.find_one({"email": data.email.lower()})
    if not user or not verify_password(data.password, user["password"]):
        raise HTTPException(status_code=401, detail="Invalid credentials")
    token = create_token(user["id"])
    return {"token": token, "user": {"id": user["id"], "name": user["name"], "email": user["email"]}}

@api_router.get("/auth/me")
async def me(user=Depends(get_current_user)):
    return user

# ---------------- Project Routes ----------------
@api_router.post("/projects")
async def create_project(data: ProjectCreate, user=Depends(get_current_user)):
    pid = str(uuid.uuid4())
    doc = {
        "id": pid,
        "user_id": user["id"],
        "name": data.name,
        "description": data.description or "",
        "audio_filename": None,
        "transcript": "",
        "requirements": None,
        "schema_spec": None,
        "api_spec": None,
        "wireframe_html": "",
        "status": "draft",
        "created_at": now_iso(),
        "updated_at": now_iso(),
    }
    await db.projects.insert_one(doc)
    doc.pop("_id", None)
    return doc

@api_router.get("/projects")
async def list_projects(user=Depends(get_current_user)):
    projects = await db.projects.find({"user_id": user["id"]}, {"_id": 0}).sort("updated_at", -1).to_list(500)
    return projects

@api_router.get("/projects/{project_id}")
async def get_project(project_id: str, user=Depends(get_current_user)):
    p = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
    if not p:
        raise HTTPException(status_code=404, detail="Project not found")
    return p

@api_router.patch("/projects/{project_id}")
async def update_project(project_id: str, data: ProjectUpdate, user=Depends(get_current_user)):
    update = {k: v for k, v in data.model_dump(exclude_none=True).items()}
    if not update:
        raise HTTPException(status_code=400, detail="No fields to update")
    update["updated_at"] = now_iso()
    res = await db.projects.update_one({"id": project_id, "user_id": user["id"]}, {"$set": update})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Project not found")
    p = await db.projects.find_one({"id": project_id}, {"_id": 0})
    return p

@api_router.delete("/projects/{project_id}")
async def delete_project(project_id: str, user=Depends(get_current_user)):
    res = await db.projects.delete_one({"id": project_id, "user_id": user["id"]})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Project not found")
    return {"ok": True}

# ---------------- Audio upload & Transcription (Whisper) ----------------
@api_router.post("/projects/{project_id}/transcribe")
async def transcribe_audio(project_id: str, file: UploadFile = File(...), user=Depends(get_current_user)):
    p = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not p:
        raise HTTPException(status_code=404, detail="Project not found")
    ext = (file.filename.split(".")[-1] if "." in file.filename else "mp3").lower()
    if ext not in {"mp3", "mp4", "mpeg", "mpga", "m4a", "wav", "webm"}:
        raise HTTPException(status_code=400, detail=f"Unsupported audio format: {ext}")
    fname = f"{project_id}_{uuid.uuid4().hex}.{ext}"
    fpath = UPLOAD_DIR / fname
    content = await file.read()
    if len(content) > 25 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Audio file exceeds 25MB limit")
    fpath.write_bytes(content)

    try:
        # Whisper transcription is CPU/GPU-bound and blocking, run it in a thread
        # so it doesn't block the FastAPI event loop.
        import asyncio
        result = await asyncio.to_thread(whisper_model.transcribe, str(fpath))
        transcript_text = result.get("text", "").strip()
    except Exception as e:
        logger.exception("Transcription failed")
        raise HTTPException(status_code=500, detail=f"Transcription failed: {str(e)}")

    await db.projects.update_one(
        {"id": project_id},
        {"$set": {"audio_filename": fname, "transcript": transcript_text, "status": "transcribed", "updated_at": now_iso()}},
    )
    return {"transcript": transcript_text}

@api_router.post("/projects/{project_id}/transcript")
async def save_transcript(project_id: str, data: TranscriptUpdate, user=Depends(get_current_user)):
    res = await db.projects.update_one(
        {"id": project_id, "user_id": user["id"]},
        {"$set": {"transcript": data.transcript, "updated_at": now_iso()}},
    )
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Project not found")
    return {"ok": True}

# ---------------- AI Generation ----------------
REQ_SYSTEM = "You are a senior software business analyst. Extract structured requirements from meeting transcripts. Return ONLY valid JSON, no markdown, no commentary."
SCHEMA_SYSTEM = "You are a senior database architect. Suggest a clean relational schema. Return ONLY valid JSON."
API_SYSTEM = "You are a senior API designer. Suggest REST endpoints. Return ONLY valid JSON."
WIRE_SYSTEM = "You are a senior UI designer creating low-fidelity HTML wireframes using Tailwind CSS. Output ONLY raw HTML (no <html> or <head>), boxed layouts with placeholders, dark theme friendly."

@api_router.post("/projects/{project_id}/extract")
async def extract_requirements(project_id: str, user=Depends(get_current_user)):
    p = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not p:
        raise HTTPException(status_code=404, detail="Project not found")
    transcript = p.get("transcript", "")
    if not transcript.strip():
        raise HTTPException(status_code=400, detail="Transcript is empty")

    prompt = f"""Read the following meeting transcript and extract structured software requirements.
Return ONLY valid JSON with this exact structure:
{{
  "project_title": "string",
  "summary": "1-2 sentence summary",
  "functional_requirements": [{{"id": "FR1", "title": "string", "description": "string", "priority": "High|Medium|Low"}}],
  "non_functional_requirements": [{{"id": "NFR1", "title": "string", "description": "string", "priority": "High|Medium|Low"}}],
  "user_roles": ["string"],
  "core_features": ["string"],
  "constraints": ["string"],
  "assumptions": ["string"]
}}

Transcript:
{transcript}"""

    try:
        data = await gemini_json(REQ_SYSTEM, prompt)
    except Exception as e:
        logger.exception("Extraction failed")
        raise HTTPException(status_code=500, detail=f"Extraction failed: {str(e)}")

    await db.projects.update_one(
        {"id": project_id},
        {"$set": {"requirements": data, "status": "extracted", "updated_at": now_iso()}},
    )
    return data

@api_router.post("/projects/{project_id}/schema")
async def generate_schema(project_id: str, user=Depends(get_current_user)):
    p = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not p or not p.get("requirements"):
        raise HTTPException(status_code=400, detail="Extract requirements first")
    prompt = f"""Based on the requirements below, suggest a clean relational DB schema.
Return ONLY valid JSON:
{{
  "tables": [
    {{
      "table_name": "string",
      "description": "string",
      "columns": [{{"name": "string", "type": "string", "primary_key": false, "nullable": true, "description": "string"}}],
      "relationships": [{{"references": "table.column", "type": "one-to-many|many-to-one|one-to-one"}}]
    }}
  ]
}}

Requirements: {json.dumps(p['requirements'])}"""
    try:
        data = await gemini_json(SCHEMA_SYSTEM, prompt)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Schema gen failed: {str(e)}")
    await db.projects.update_one({"id": project_id}, {"$set": {"schema_spec": data, "updated_at": now_iso()}})
    return data

@api_router.post("/projects/{project_id}/api-spec")
async def generate_api_spec(project_id: str, user=Depends(get_current_user)):
    p = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not p or not p.get("requirements"):
        raise HTTPException(status_code=400, detail="Extract requirements first")
    prompt = f"""Based on the requirements, suggest REST API endpoints.
Return ONLY valid JSON:
{{
  "base_url": "/api",
  "endpoints": [
    {{"method": "GET|POST|PUT|PATCH|DELETE", "path": "/resource", "description": "string", "auth_required": true, "request_body": "string", "response": "string"}}
  ]
}}

Requirements: {json.dumps(p['requirements'])}"""
    try:
        data = await gemini_json(API_SYSTEM, prompt)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"API gen failed: {str(e)}")
    await db.projects.update_one({"id": project_id}, {"$set": {"api_spec": data, "updated_at": now_iso()}})
    return data

@api_router.post("/projects/{project_id}/wireframe")
async def generate_wireframe(project_id: str, user=Depends(get_current_user)):
    p = await db.projects.find_one({"id": project_id, "user_id": user["id"]})
    if not p or not p.get("requirements"):
        raise HTTPException(status_code=400, detail="Extract requirements first")
    reqs = p["requirements"]
    core = reqs.get("core_features", [])
    roles = reqs.get("user_roles", [])
    title = reqs.get("project_title", "App")

    prompt = f"""Generate a low-fidelity HTML wireframe for the following application using Tailwind CSS.
- Project: {title}
- Core features: {core}
- User roles: {roles}
- Output ONLY the inner HTML (no <html>, <head>, <body> tags).
- Use a top navbar, sidebar with role-based menu items, and a main content area showing key screens as stacked sections.
- Use placeholder boxes (border-dashed), grey rectangles for images, sample text for labels.
- Dark theme: use bg-gray-900, text-gray-100, border-gray-700, etc.
- Keep it visually clean and easy to scan."""
    try:
        html = await gemini_text(WIRE_SYSTEM, prompt)
        # strip code fences if present
        html = re.sub(r"^```(?:html)?\n?", "", html.strip()).rstrip("`").rstrip("\n").strip()
        if html.endswith("```"):
            html = html[:-3].strip()
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Wireframe gen failed: {str(e)}")

    await db.projects.update_one({"id": project_id}, {"$set": {"wireframe_html": html, "updated_at": now_iso()}})
    return {"wireframe_html": html}

# ---------------- SRS DOCX Export ----------------
@api_router.get("/projects/{project_id}/srs.docx")
async def export_srs_docx(project_id: str, token: str):
    # token via query param so browser <a> link works
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
        user_id = payload.get("sub")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")
    p = await db.projects.find_one({"id": project_id, "user_id": user_id})
    if not p:
        raise HTTPException(status_code=404, detail="Project not found")
    reqs = p.get("requirements") or {}
    if not reqs:
        raise HTTPException(status_code=400, detail="No requirements to export")

    doc = Document()
    title = doc.add_heading(reqs.get("project_title", p["name"]), level=0)
    doc.add_paragraph(f"Software Requirements Specification — Generated by ProtoPilot")
    doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(reqs.get("summary", p.get("description", "")))

    doc.add_heading("2. Scope", level=1)
    doc.add_paragraph("Core Features:")
    for f in reqs.get("core_features", []):
        doc.add_paragraph(f, style="List Bullet")

    doc.add_heading("3. User Roles", level=1)
    for r in reqs.get("user_roles", []):
        doc.add_paragraph(r, style="List Bullet")

    doc.add_heading("4. Functional Requirements", level=1)
    for fr in reqs.get("functional_requirements", []):
        h = doc.add_paragraph()
        run = h.add_run(f"{fr.get('id','FR')}: {fr.get('title','')}")
        run.bold = True
        doc.add_paragraph(f"Priority: {fr.get('priority','Medium')}")
        doc.add_paragraph(fr.get("description", ""))

    doc.add_heading("5. Non-Functional Requirements", level=1)
    for nfr in reqs.get("non_functional_requirements", []):
        h = doc.add_paragraph()
        run = h.add_run(f"{nfr.get('id','NFR')}: {nfr.get('title','')}")
        run.bold = True
        doc.add_paragraph(f"Priority: {nfr.get('priority','Medium')}")
        doc.add_paragraph(nfr.get("description", ""))

    doc.add_heading("6. Assumptions", level=1)
    for a in reqs.get("assumptions", []):
        doc.add_paragraph(a, style="List Bullet")

    doc.add_heading("7. Constraints", level=1)
    for c in reqs.get("constraints", []):
        doc.add_paragraph(c, style="List Bullet")

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    fname = re.sub(r"[^a-zA-Z0-9_-]", "_", reqs.get("project_title", p["name"]))[:40] + "_SRS.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )

@api_router.get("/projects/{project_id}/wireframe.html")
async def export_wireframe_html(project_id: str, token: str):
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
        user_id = payload.get("sub")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")
    p = await db.projects.find_one({"id": project_id, "user_id": user_id})
    if not p or not p.get("wireframe_html"):
        raise HTTPException(status_code=404, detail="No wireframe")
    html_doc = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{p['name']} - Wireframe</title>
<script src="https://cdn.tailwindcss.com"></script></head>
<body class="bg-gray-900 text-gray-100 min-h-screen">{p['wireframe_html']}</body></html>"""
    buf = BytesIO(html_doc.encode())
    fname = re.sub(r"[^a-zA-Z0-9_-]", "_", p["name"])[:40] + "_wireframe.html"
    return StreamingResponse(buf, media_type="text/html",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})

@api_router.get("/")
async def root():
    return {"message": "ProtoPilot API"}

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', 'http://localhost:3000').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()