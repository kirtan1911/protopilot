from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import Response
from database import db
from deps import get_current_user
import json
import io
import zipfile
from docx import Document

router = APIRouter(prefix="/projects", tags=["Exports"])

def wrap_wireframe_html(page_name: str, html_content: str, manifest: dict = None) -> str:
    """Wraps raw AI HTML snippets into a complete standalone HTML5 document with Tailwind CSS, Lucide icons, Inter fonts, and custom styling."""
    if not html_content:
        html_content = "<div class='p-8 text-center text-gray-400'>No content available.</div>"

    # If it's already a full HTML document, return as is
    if "<!doctype html>" in html_content.lower() or "<html" in html_content.lower():
        return html_content

    design_system = manifest.get("design_system", {}) if isinstance(manifest, dict) else {}
    colors_json = json.dumps(design_system.get("colors", {
        "primary": "#6366f1",
        "secondary": "#a855f7",
        "background": "#0d1117",
        "surface": "#161b22"
    }))

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{page_name} - Wireframe Prototype</title>
    <!-- Tailwind CSS CDN -->
    <script src="https://cdn.tailwindcss.com"></script>
    <!-- Custom Tailwind Configuration -->
    <script>
        tailwind.config = {{
            darkMode: 'class',
            theme: {{
                extend: {{
                    colors: {colors_json},
                    fontFamily: {{ sans: ['Inter', 'sans-serif'] }}
                }}
            }}
        }}
    </script>
    <!-- Google Fonts & Lucide Icons -->
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap" rel="stylesheet">
    <script src="https://unpkg.com/lucide@latest"></script>
    <link rel="stylesheet" href="styles.css">
    <style>
        body {{
            margin: 0;
            padding: 0;
            background-color: #0d1117;
            color: #e5e7eb;
            font-family: 'Inter', sans-serif;
        }}
        .glass {{
            background: rgba(255, 255, 255, 0.03);
            backdrop-filter: blur(10px);
            -webkit-backdrop-filter: blur(10px);
            border: 1px solid rgba(255, 255, 255, 0.05);
        }}
        .glow-indigo {{
            box-shadow: 0 0 20px rgba(99, 102, 241, 0.35);
        }}
        ::-webkit-scrollbar {{ width: 8px; height: 8px; }}
        ::-webkit-scrollbar-track {{ background: #0d1117; }}
        ::-webkit-scrollbar-thumb {{ background: #21262d; border-radius: 4px; }}
        ::-webkit-scrollbar-thumb:hover {{ background: #30363d; }}
    </style>
</head>
<body class="bg-[#0d1117] text-gray-100 min-h-screen">
{html_content}
    <script>
        document.addEventListener('DOMContentLoaded', () => {{
            if (window.lucide) {{
                lucide.createIcons();
            }}
            // Link forms/buttons to prevent accidental page resets
            document.querySelectorAll('form').forEach(f => {{
                f.addEventListener('submit', (e) => e.preventDefault());
            }});
        }});
    </script>
</body>
</html>"""

@router.get("/{project_id}/export")
async def export_project(project_id: str, user=Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    return project

@router.get("/{project_id}/srs.docx")
async def export_srs_docx(project_id: str, user=Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
        
    reqs = project.get("requirements", {})
    
    doc = Document()
    doc.add_heading(reqs.get("project_title", project.get("name", "Software Requirements Specification")), 0)
    
    if reqs.get("summary"):
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(reqs.get("summary"))
        
    if reqs.get("user_roles"):
        doc.add_heading("User Roles & Personas", level=1)
        for role in reqs.get("user_roles", []):
            doc.add_paragraph(str(role), style='List Bullet')
            
    if reqs.get("core_features"):
        doc.add_heading("Core Features", level=1)
        for feat in reqs.get("core_features", []):
            doc.add_paragraph(str(feat), style='List Bullet')
            
    if reqs.get("functional_requirements"):
        doc.add_heading("Functional Requirements", level=1)
        for fr in reqs.get("functional_requirements", []):
            title = fr.get("title", "Requirement") if isinstance(fr, dict) else str(fr)
            desc = fr.get("description", "") if isinstance(fr, dict) else ""
            pri = fr.get("priority", "Medium") if isinstance(fr, dict) else "Medium"
            doc.add_heading(f"[{pri}] {title}", level=2)
            if desc:
                doc.add_paragraph(desc)

    if reqs.get("non_functional_requirements"):
        doc.add_heading("Non-Functional Requirements", level=1)
        for nfr in reqs.get("non_functional_requirements", []):
            title = nfr.get("title", "Requirement") if isinstance(nfr, dict) else str(nfr)
            desc = nfr.get("description", "") if isinstance(nfr, dict) else ""
            doc.add_heading(title, level=2)
            if desc:
                doc.add_paragraph(desc)

    if reqs.get("constraints"):
        doc.add_heading("System Constraints", level=1)
        for c in reqs.get("constraints", []):
            doc.add_paragraph(str(c), style='List Bullet')

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    filename = f"{project.get('name', 'project').replace(' ', '_')}_SRS.docx"
    return Response(
        content=file_stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/{project_id}/api_spec.json")
async def export_api_spec(project_id: str, user=Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
        
    api_spec = project.get("api_spec", {"info": {"title": project.get("name"), "version": "1.0.0"}, "endpoints": []})
    formatted_json = json.dumps(api_spec, indent=2)
    
    filename = f"{project.get('name', 'project').replace(' ', '_')}_api_spec.json"
    return Response(
        content=formatted_json,
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/{project_id}/schema.sql")
async def export_schema_sql(project_id: str, user=Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
        
    schema_spec = project.get("schema_spec", {})
    sql_lines = [f"-- SQL Schema for {project.get('name')}", "-- Generated by ProtoPilot\n"]
    
    tables = schema_spec.get("tables", []) if isinstance(schema_spec, dict) else []
    for table in tables:
        tname = table.get("name", "table_name")
        sql_lines.append(f"CREATE TABLE {tname} (")
        col_defs = []
        for col in table.get("columns", []):
            cname = col.get("name", "id")
            ctype = col.get("type", "VARCHAR(255)")
            nullable = "" if col.get("nullable", True) else " NOT NULL"
            pk = " PRIMARY KEY" if col.get("primary_key") else ""
            col_defs.append(f"    {cname} {ctype}{nullable}{pk}")
        sql_lines.append(",\n".join(col_defs))
        sql_lines.append(");\n")
        
    if len(sql_lines) <= 2:
        sql_lines.append("-- Default SQL Schema Structure")
        sql_lines.append("CREATE TABLE users (id VARCHAR(36) PRIMARY KEY, email VARCHAR(255) UNIQUE, created_at TIMESTAMP);")
        sql_lines.append("CREATE TABLE projects (id VARCHAR(36) PRIMARY KEY, name VARCHAR(255), user_id VARCHAR(36));")
        
    sql_content = "\n".join(sql_lines)
    filename = f"{project.get('name', 'project').replace(' ', '_')}_schema.sql"
    return Response(
        content=sql_content,
        media_type="text/plain",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/{project_id}/wireframes.zip")
async def export_wireframes_zip(project_id: str, user=Depends(get_current_user)):
    project = await db.projects.find_one({"id": project_id, "user_id": user["id"]}, {"_id": 0})
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
        
    pages = project.get("wireframe_pages", {})
    manifest = project.get("wireframe_manifest", {})
    manifest_pages = manifest.get("pages", []) if isinstance(manifest, dict) else []
    
    page_names = {p.get("id"): p.get("name", "Page") for p in manifest_pages} if isinstance(manifest_pages, list) else {}

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
        # 1. Write manifest JSON
        zip_file.writestr("manifest.json", json.dumps(manifest, indent=2))
        
        # 2. Write custom CSS stylesheet
        zip_file.writestr("styles.css", """/* ProtoPilot Standalone Wireframe Styles */
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');
body { margin: 0; padding: 0; background-color: #0d1117; color: #e5e7eb; font-family: 'Inter', sans-serif; }
.glass { background: rgba(255, 255, 255, 0.03); backdrop-filter: blur(10px); -webkit-backdrop-filter: blur(10px); border: 1px solid rgba(255, 255, 255, 0.05); }
.glow-indigo { box-shadow: 0 0 20px rgba(99, 102, 241, 0.35); }
""")

        # 3. Write HTML pages
        if pages:
            first_page_html = None
            has_index = False
            
            for page_id, raw_html in pages.items():
                pname = page_names.get(page_id, page_id.replace("_", " ").title())
                full_html = wrap_wireframe_html(pname, raw_html, manifest)
                zip_file.writestr(f"{page_id}.html", full_html)
                
                if page_id in ("index", "landing", "home"):
                    has_index = True
                if first_page_html is None:
                    first_page_html = full_html

            if not has_index and first_page_html:
                zip_file.writestr("index.html", first_page_html)
        else:
            default_html = wrap_wireframe_html(
                "Home",
                '<div class="flex flex-col items-center justify-center min-h-screen p-8 text-center"><h1 class="text-3xl font-bold text-indigo-400 mb-4">ProtoPilot Prototype</h1><p class="text-gray-400 max-w-md">No wireframe pages have been generated yet. Open ProtoPilot and click "Generate Prototype".</p></div>',
                manifest
            )
            zip_file.writestr("index.html", default_html)
            
    zip_buffer.seek(0)
    filename = f"{project.get('name', 'project').replace(' ', '_')}_wireframes.zip"
    return Response(
        content=zip_buffer.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
